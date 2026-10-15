"""模板版本 + 属主/授权（ACL）API 集成测试。"""

import io
from collections.abc import AsyncGenerator
from pathlib import Path
from types import SimpleNamespace

import pytest_asyncio
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import (
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)

from app.api.deps import get_current_user
from app.config import get_settings
from app.database import get_db
from app.main import app
from app.models import Base
from app.models.template import TemplateVersion, WordTemplate
from app.models.user import User
from app.services import onlyoffice_service


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest_asyncio.fixture
async def ctx(tmp_path, monkeypatch, fake_user: User):
    monkeypatch.setattr(get_settings(), "storage_dir", str(tmp_path / "storage"))
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(
        engine, class_=AsyncSession, expire_on_commit=False
    )

    async def _get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as s:
            yield s

    app.dependency_overrides[get_db] = _get_db
    holder = {"user": fake_user}
    app.dependency_overrides[get_current_user] = lambda: holder["user"]
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield SimpleNamespace(
            client=client,
            session_factory=session_factory,
            admin=fake_user,
            set_user=lambda u: holder.__setitem__("user", u),
        )
    app.dependency_overrides.clear()
    await engine.dispose()


async def _add_user(session_factory, uid: int, username: str, role: str = "user") -> User:
    user = User(
        id=uid,
        username=username,
        email=f"{username}@x.com",
        display_name=username,
        source="local",
        role=role,
        is_active=True,
    )
    async with session_factory() as s:
        s.add(user)
        await s.commit()
        s.expunge(user)
    return user


async def _upload(client: AsyncClient, name: str = "月报模板") -> dict:
    resp = await client.post(
        "/api/templates",
        data={"name": name},
        files={
            "file": (
                "tpl.docx",
                make_docx_bytes(["CPU: {{ cpu }}", "摘要: {{ summary }}"]),
                "application/octet-stream",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


class TestUploadVersion:
    async def test_upload_creates_owner_and_version1(self, ctx) -> None:
        body = await _upload(ctx.client)
        assert body["owner_id"] == 1
        assert body["permission"] == "manage"
        assert body["current_version"]["version_no"] == 1
        assert body["current_version"]["is_current"] is True
        assert [p["name"] for p in body["placeholders"]] == ["cpu", "summary"]


class TestAcl:
    async def test_grant_flow(self, ctx) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        u2 = await _add_user(ctx.session_factory, 2, "user2")

        ctx.set_user(u2)
        assert (await ctx.client.get("/api/templates")).json() == []
        assert (await ctx.client.get(f"/api/templates/{tpl_id}")).status_code == 403

        ctx.set_user(ctx.admin)
        r = await ctx.client.put(
            f"/api/templates/{tpl_id}/grants",
            json={"grants": [{"user_id": 2, "permission": "view"}]},
        )
        assert r.status_code == 200 and r.json()[0]["permission"] == "view"

        ctx.set_user(u2)
        assert (await ctx.client.get(f"/api/templates/{tpl_id}")).json()["permission"] == "view"
        assert len((await ctx.client.get("/api/templates")).json()) == 1
        # view 不可绑定
        r = await ctx.client.put(
            f"/api/templates/{tpl_id}/bindings", json={"bindings": []}
        )
        assert r.status_code == 403

        ctx.set_user(ctx.admin)
        await ctx.client.put(
            f"/api/templates/{tpl_id}/grants",
            json={"grants": [{"user_id": 2, "permission": "edit"}]},
        )
        ctx.set_user(u2)
        r = await ctx.client.put(
            f"/api/templates/{tpl_id}/bindings", json={"bindings": []}
        )
        assert r.status_code == 200
        # edit 不可授权
        r = await ctx.client.put(
            f"/api/templates/{tpl_id}/grants", json={"grants": []}
        )
        assert r.status_code == 403

    async def test_transfer_owner(self, ctx) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        u2 = await _add_user(ctx.session_factory, 2, "user2")
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/transfer", json={"owner_id": 2}
        )
        assert r.status_code == 200 and r.json()["owner_id"] == 2
        ctx.set_user(u2)
        assert (await ctx.client.get(f"/api/templates/{tpl_id}")).json()["permission"] == "manage"

    async def test_grant_unknown_user_rejected(self, ctx) -> None:
        body = await _upload(ctx.client)
        r = await ctx.client.put(
            f"/api/templates/{body['id']}/grants",
            json={"grants": [{"user_id": 999, "permission": "view"}]},
        )
        assert r.status_code == 422


class TestVersions:
    async def test_list_and_activate(self, ctx) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        versions = (await ctx.client.get(f"/api/templates/{tpl_id}/versions")).json()
        assert len(versions) == 1 and versions[0]["version_no"] == 1

        async with ctx.session_factory() as s:
            v2 = TemplateVersion(
                template_id=tpl_id,
                version_no=2,
                file_path="v2.docx",
                placeholders_json=["cpu"],
                source="edit",
            )
            s.add(v2)
            await s.commit()
            await s.refresh(v2)
            v2_id = v2.id

        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/versions/{v2_id}/activate"
        )
        assert r.status_code == 200 and r.json()["is_current"] is True
        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 2


class TestFileAndEditor:
    async def test_file_requires_auth_or_token(self, ctx) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]
        assert (await ctx.client.get(f"/api/templates/{tpl_id}/file")).status_code == 401

        token = onlyoffice_service.sign_token({"t": tpl_id, "v": vid, "k": "file"})
        r = await ctx.client.get(f"/api/templates/{tpl_id}/file?token={token}")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats"
        )

    async def test_bad_token_rejected(self, ctx) -> None:
        body = await _upload(ctx.client)
        r = await ctx.client.get(f"/api/templates/{body['id']}/file?token=bad")
        assert r.status_code == 401

    async def test_editor_disabled_when_not_configured(self, ctx, monkeypatch) -> None:
        monkeypatch.setattr(get_settings(), "onlyoffice_url", "")
        monkeypatch.setattr(get_settings(), "onlyoffice_jwt_secret", "")
        monkeypatch.setattr(get_settings(), "public_base_url", "")
        body = await _upload(ctx.client)
        r = await ctx.client.post(f"/api/templates/{body['id']}/editor")
        assert r.status_code == 200
        assert r.json()["enabled"] is False


class TestCallback:
    async def test_callback_creates_new_version(self, ctx, monkeypatch) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]

        async def fake_download(url: str) -> bytes:
            return make_docx_bytes(["CPU: {{ cpu }}", "新变量: {{ extra }}"])

        monkeypatch.setattr(onlyoffice_service, "download_document", fake_download)

        token = onlyoffice_service.sign_token(
            {"t": tpl_id, "v": vid, "u": 1, "k": "cb"}
        )
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/callback?token={token}",
            json={"status": 2, "url": "http://ds/download"},
        )
        assert r.status_code == 200 and r.json() == {"error": 0}

        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 2
        assert [p["name"] for p in detail["placeholders"]] == ["cpu", "extra"]

    async def test_callback_bad_token_rejected(self, ctx) -> None:
        body = await _upload(ctx.client)
        r = await ctx.client.post(
            f"/api/templates/{body['id']}/callback?token=bad",
            json={"status": 2, "url": "http://ds/x"},
        )
        assert r.status_code == 403


class TestCallbackVersionSemantics:
    """forcesave(6) 不建版本；内容与当前版本一致(2) 时跳过，避免重复版本。"""

    async def test_force_save_creates_no_version(self, ctx, monkeypatch) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]

        called = {"n": 0}

        async def fake_download(url: str) -> bytes:
            called["n"] += 1
            return make_docx_bytes(["CPU: {{ cpu }}"])

        monkeypatch.setattr(onlyoffice_service, "download_document", fake_download)
        token = onlyoffice_service.sign_token({"t": tpl_id, "v": vid, "u": 1, "k": "cb"})
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/callback?token={token}",
            json={"status": 6, "url": "http://ds/download"},
        )
        assert r.status_code == 200 and r.json() == {"error": 0}
        assert called["n"] == 0, "强制保存不应下载文件"

        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 1

    async def test_identical_content_skips_new_version(self, ctx, monkeypatch) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]

        # 取当前版本的文件字节（直接从存储读，避免依赖 HTTP 端点）
        async with ctx.session_factory() as s:
            v = await s.scalar(
                select(TemplateVersion).where(
                    TemplateVersion.id == body["current_version"]["id"]
                )
            )
            same_bytes = Path(v.file_path).read_bytes()

        async def fake_download(url: str) -> bytes:
            return same_bytes

        monkeypatch.setattr(onlyoffice_service, "download_document", fake_download)
        token = onlyoffice_service.sign_token({"t": tpl_id, "v": vid, "u": 1, "k": "cb"})
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/callback?token={token}",
            json={"status": 2, "url": "http://ds/download"},
        )
        assert r.status_code == 200 and r.json() == {"error": 0}

        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 1, "内容未变不应新增版本"

    async def test_changed_content_creates_version(self, ctx, monkeypatch) -> None:
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]

        async def fake_download(url: str) -> bytes:
            return make_docx_bytes(["CPU: {{ cpu }}", "改了"])

        monkeypatch.setattr(onlyoffice_service, "download_document", fake_download)
        token = onlyoffice_service.sign_token({"t": tpl_id, "v": vid, "u": 1, "k": "cb"})
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/callback?token={token}",
            json={"status": 2, "url": "http://ds/download"},
        )
        assert r.json() == {"error": 0}
        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 2

    async def test_repeated_force_save_then_one_save(self, ctx, monkeypatch) -> None:
        """模拟真实编辑：多次强制保存 + 一次最终保存 -> 只多一个版本。"""
        body = await _upload(ctx.client)
        tpl_id = body["id"]
        vid = body["current_version"]["id"]
        token = onlyoffice_service.sign_token({"t": tpl_id, "v": vid, "u": 1, "k": "cb"})

        async def fake_download(url: str) -> bytes:
            return make_docx_bytes(["CPU: {{ cpu }}", f"编辑 {url}"])

        monkeypatch.setattr(onlyoffice_service, "download_document", fake_download)
        for i in range(3):
            r = await ctx.client.post(
                f"/api/templates/{tpl_id}/callback?token={token}",
                json={"status": 6, "url": f"http://ds/fs{i}"},
            )
            assert r.json() == {"error": 0}
        r = await ctx.client.post(
            f"/api/templates/{tpl_id}/callback?token={token}",
            json={"status": 2, "url": "http://ds/final"},
        )
        assert r.json() == {"error": 0}
        detail = (await ctx.client.get(f"/api/templates/{tpl_id}")).json()
        assert detail["current_version"]["version_no"] == 2
