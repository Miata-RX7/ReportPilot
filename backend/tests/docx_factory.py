"""测试用 docx 工厂：动态生成模板，无二进制 fixture 入库。

可精确控制 run 拆分（模拟 Word 把 {{var}} 拆成多段）与表格/页眉场景。
"""

from pathlib import Path

from docx import Document


def make_docx(
    path: Path,
    paragraphs: list[str] | None = None,
    table_cells: list[str] | None = None,
    split_runs: list[list[str]] | None = None,
    header_texts: list[str] | None = None,
) -> Path:
    doc = Document()
    for text in paragraphs or []:
        doc.add_paragraph(text)
    for parts in split_runs or []:
        p = doc.add_paragraph()
        for part in parts:
            p.add_run(part)
    if table_cells:
        table = doc.add_table(rows=len(table_cells), cols=1)
        for i, cell_text in enumerate(table_cells):
            table.rows[i].cells[0].text = cell_text
    for text in header_texts or []:
        doc.sections[0].header.add_paragraph(text)
    doc.save(str(path))
    return path


def docx_bytes(**kwargs: object) -> bytes:
    """生成 docx 并返回字节流（API 上传测试用）。"""
    import io

    doc = Document()
    for text in kwargs.get("paragraphs", []) or []:  # type: ignore[union-attr]
        doc.add_paragraph(text)
    for parts in kwargs.get("split_runs", []) or []:  # type: ignore[union-attr]
        p = doc.add_paragraph()
        for part in parts:
            p.add_run(part)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def read_docx_text(path: Path) -> str:
    """读回 docx 全部正文+表格文本（断言替换结果用）。"""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
