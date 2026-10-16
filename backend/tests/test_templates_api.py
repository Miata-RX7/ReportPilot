"""模板 API 集成测试（SQLite 内存库 + ASGITransport + tmp storage）。"""

import io
from collections.abc import AsyncGenerator
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pytest
import pytest_asyncio
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import (
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)

import app.api.deps as deps_module
from app.api.deps import get_current_user
from app.config import get_settings
from app.database import get_db
from app.main import app
from app.models import Base
from app.models.user import User
from app.services.adapters.base import DataSourceAdapter
from app.services.preview_service import AI_SUMMARY_PLACEHOLDER

SEP = int(datetime(2026, 9, 1, tzinfo=timezone.utc).timestamp())

DATASOURCE_PAYLOAD = {
    "name": "生产 Zabbix",
    "type": "zabbix",
    "config": {"url": "http://zabbix.example.com", "auth_type": "bearer", "timeout": 10, "verify_ssl": False},
    "secret": "zabbix-token-xyz7890",
}

METRIC_PAYLOAD = {
    "metric_type": "scalar",
    "name": "cpu_usage",
    "data_source_id": 1,
    "query_config": {"host_group_ids": ["4"], "host_ids": ["1"], "item_ids": ["101"]},
    "calc_formula": {"agg": "avg", "compare": "none"},
    "unit": "%",
    "period_type": "monthly",
}


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FakeAdapter(DataSourceAdapter):
    async def test_connection(self) -> bool:
        return True

    async def list_host_groups(self) -> list[dict[str, Any]]:
        return [{"group_id": 4, "name": "Linux servers"}]

    async def list_hosts(self, group_ids: list[str]) -> list[dict[str, Any]]:
        return [
            {
                "host_id": 1,
                "name": "web-01",
                "groups": [{"group_id": 4, "name": "Linux servers"}],
            }
        ]

    async def list_items(
        self, host_ids: list[str], value_types: list[int] | None = None
    ) -> list[dict[str, Any]]:
        return [
            {
                "item_id": 101,
                "name": "cpu",
                "key": "system.cpu.util",
                "host_id": 1,
                "host": "web-01",
                "value_type": 0,
                "units": "%",
            }
        ]

    async def fetch_latest_value(  # type: ignore[no-untyped-def]
        self, item_id: int, time_from: int | None = None, time_till: int | None = None
    ):
        return None


    async def get_item_meta(self, item_id: int):
        return None
    async def fetch_trend_data(
        self, item_ids: list[int], time_from: int, time_till: int
    ) -> list[dict[str, Any]]:
        return [
            {
                "item_id": 101,
                "clock": SEP + 3600,
                "num": 60,
                "value_min": 30.0,
                "value_avg": 30.0,
                "value_max": 30.0,
            }
        ]


@pytest_asyncio.fixture
async def api(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fake_user: User
) -> AsyncGenerator[AsyncClient, None]:
    monkeypatch.setattr(get_settings(), "storage_dir", str(tmp_path / "storage"))
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(
        engine, class_=AsyncSession, expire_on_commit=False
    )

    async def _get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = _get_db
    app.dependency_overrides[get_current_user] = lambda: fake_user
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield client
    app.dependency_overrides.clear()
    await engine.dispose()


async def _upload(client: AsyncClient, paragraphs: list[str], name: str = "月报模板") -> dict:
    resp = await client.post(
        "/api/templates",
        data={"name": name},
        files={"file": ("tpl.docx", make_docx_bytes(paragraphs), "application/octet-stream")},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


class TestUpload:
    async def test_upload_list_detail(self, api: AsyncClient) -> None:
        body = await _upload(api, ["CPU: {{ cpu }}", "摘要: {{ summary }}"])
        # 上传响应与详情都返回带独占标记的结构
        assert [p["name"] for p in body["placeholders"]] == ["cpu", "summary"]
        assert [p["is_alone_in_paragraph"] for p in body["placeholders"]] == [False, False]
        assert body["unbound"] == ["cpu", "summary"]

        list_resp = await api.get("/api/templates")
        item = list_resp.json()[0]
        assert item["placeholder_count"] == 2
        assert item["bound_count"] == 0

        detail = await api.get(f"/api/templates/{body['id']}")
        assert detail.status_code == 200

    async def test_violation_rejected(self, api: AsyncClient) -> None:
        resp = await api.post(
            "/api/templates",
            data={"name": "坏模板"},
            files={"file": ("tpl.docx", make_docx_bytes(["{{ a.b }}"]), "application/octet-stream")},
        )
        assert resp.status_code == 422
        assert "a.b" in resp.text

    async def test_non_docx_extension_rejected(self, api: AsyncClient) -> None:
        resp = await api.post(
            "/api/templates",
            data={"name": "x"},
            files={"file": ("tpl.doc", b"PK\x03\x04fake", "application/octet-stream")},
        )
        assert resp.status_code == 422

    async def test_bad_magic_rejected(self, api: AsyncClient) -> None:
        resp = await api.post(
            "/api/templates",
            data={"name": "x"},
            files={"file": ("tpl.docx", b"not-a-zip-file", "application/octet-stream")},
        )
        assert resp.status_code == 422
        assert "magic" in resp.json()["detail"]

    async def test_oversized_rejected(self, api: AsyncClient) -> None:
        big = b"PK\x03\x04" + b"x" * (10 * 1024 * 1024 + 1)
        resp = await api.post(
            "/api/templates",
            data={"name": "x"},
            files={"file": ("tpl.docx", big, "application/octet-stream")},
        )
        assert resp.status_code == 413

    async def test_duplicate_name_conflict(self, api: AsyncClient) -> None:
        await _upload(api, ["{{ a }}"])
        resp = await api.post(
            "/api/templates",
            data={"name": "月报模板"},
            files={"file": ("tpl.docx", make_docx_bytes(["{{ a }}"]), "application/octet-stream")},
        )
        assert resp.status_code == 409

    async def test_path_traversal_filename_sanitized(self, api: AsyncClient) -> None:
        resp = await api.post(
            "/api/templates",
            data={"name": "清洗测试"},
            files={"file": ("../../etc/passwd.docx", make_docx_bytes(["{{ a }}"]), "application/octet-stream")},
        )
        assert resp.status_code == 201


class TestBindings:
    async def _setup_metric(self, client: AsyncClient) -> int:
        await client.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        resp = await client.post("/api/metrics", json=METRIC_PAYLOAD)
        assert resp.status_code == 201, resp.text
        return resp.json()["id"]

    async def test_put_replace_and_clear(self, api: AsyncClient) -> None:
        metric_id = await self._setup_metric(api)
        tpl = await _upload(api, ["{{ cpu }}", "{{ summary }}"])

        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                    {"placeholder": "summary", "binding_type": "insight"},
                ]
            },
        )
        assert resp.status_code == 200, resp.text
        result = resp.json()
        assert result["cleared_count"] == 0
        assert len(result["bindings"]) == 2
        insight = [b for b in result["bindings"] if b["binding_type"] == "insight"][0]
        assert insight["binding_key"] == "summary"
        assert insight["binding_ref_id"] is None

        detail = await api.get(f"/api/templates/{tpl['id']}")
        assert detail.json()["unbound"] == []

        # 再次 PUT = 全量替换（重复绑定覆盖）
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={"bindings": [{"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id}]},
        )
        assert resp.json()["cleared_count"] == 2
        assert len(resp.json()["bindings"]) == 1

        # 空数组 = 清空
        resp = await api.put(f"/api/templates/{tpl['id']}/bindings", json={"bindings": []})
        assert resp.json()["cleared_count"] == 1
        assert resp.json()["bindings"] == []

    async def test_bind_unknown_placeholder(self, api: AsyncClient) -> None:
        tpl = await _upload(api, ["{{ cpu }}"])
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={"bindings": [{"placeholder": "ghost", "binding_type": "insight"}]},
        )
        assert resp.status_code == 422

    async def test_bind_missing_metric(self, api: AsyncClient) -> None:
        tpl = await _upload(api, ["{{ cpu }}"])
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={"bindings": [{"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": 999}]},
        )
        assert resp.status_code == 422

    async def test_metric_binding_requires_ref_id(self, api: AsyncClient) -> None:
        tpl = await _upload(api, ["{{ cpu }}"])
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={"bindings": [{"placeholder": "cpu", "binding_type": "metric"}]},
        )
        assert resp.status_code == 422

    async def test_bindings_replace_no_unique_violation(self, api: AsyncClient) -> None:
        """回归：全量替换（删旧插新同 placeholder）不得触发唯一约束冲突。"""
        metric_id = await self._setup_metric(api)
        tpl = await _upload(api, ["{{ cpu }}"])
        payload = {
            "bindings": [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id}
            ]
        }
        for _ in range(3):  # 同一 placeholder 反复覆盖
            resp = await api.put(f"/api/templates/{tpl['id']}/bindings", json=payload)
            assert resp.status_code == 200, resp.text
            assert len(resp.json()["bindings"]) == 1


class TestPreview:
    async def test_preview_end_to_end(
        self, api: AsyncClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None))
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        metric = await api.post("/api/metrics", json=METRIC_PAYLOAD)
        tpl = await _upload(api, ["CPU 使用率: {{ cpu }}", "{{ summary }}"])
        await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric.json()["id"]},
                    {"placeholder": "summary", "binding_type": "insight"},
                ]
            },
        )
        resp = await api.post(
            f"/api/templates/{tpl['id']}/preview", json={"period": "2026-09"}
        )
        assert resp.status_code == 200, resp.text
        assert resp.content[:4] == b"PK\x03\x04"
        out = tmp_path / "preview.docx"
        out.write_bytes(resp.content)
        doc = Document(str(out))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "CPU 使用率: 30" in text
        assert AI_SUMMARY_PLACEHOLDER in text

    async def test_preview_unbound_variable(
        self, api: AsyncClient, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None))
        tpl = await _upload(api, ["{{ cpu }}"])
        resp = await api.post(
            f"/api/templates/{tpl['id']}/preview", json={"period": "2026-09"}
        )
        assert resp.status_code == 422
        assert "cpu" in resp.json()["detail"]


class TestDelete:
    async def test_soft_delete(self, api: AsyncClient) -> None:
        tpl = await _upload(api, ["{{ a }}"])
        resp = await api.delete(f"/api/templates/{tpl['id']}")
        assert resp.status_code == 204
        assert (await api.get("/api/templates")).json() == []
        with_inactive = await api.get("/api/templates?include_inactive=true")
        assert len(with_inactive.json()) == 1


class TestChartTypeBindingGuard:
    """阶段 3A：curve / pie 可绑定，但要求变量在模板中独占一个段落。"""

    @staticmethod
    def _curve(api: AsyncClient, name: str = "traffic_curve") -> dict:
        return {
            **METRIC_PAYLOAD,
            "name": name,
            "metric_type": "curve",
            "query_config": {
                "host_group_ids": ["4"],
                "host_ids": ["1"],
                "series": [{"item_id": "101", "display_name": "in"}],
            },
        }

    async def test_curve_metric_allowed_when_alone(self, api: AsyncClient) -> None:
        """独占段落 -> 绑定成功（护栏已放开）。"""
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["{{ curve_x }}"])
        metric = (await api.post("/api/metrics", json=self._curve(api))).json()

        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {
                        "placeholder": "curve_x",
                        "binding_type": "metric",
                        "binding_ref_id": metric["id"],
                    }
                ]
            },
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["bindings"][0]["binding_ref_id"] == metric["id"]

    async def test_curve_metric_rejected_when_mixed_with_text(
        self, api: AsyncClient
    ) -> None:
        """与文字混排 -> 422 且提示明确（独占段落要求）。"""
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["流量：{{ curve_x }}"])
        metric = (await api.post("/api/metrics", json=self._curve(api))).json()

        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {
                        "placeholder": "curve_x",
                        "binding_type": "metric",
                        "binding_ref_id": metric["id"],
                    }
                ]
            },
        )
        assert resp.status_code == 422
        assert "独占一段" in resp.text
        assert "curve_x" in resp.text

    async def test_pie_metric_rejected_when_mixed(self, api: AsyncClient) -> None:
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["占比 {{ pie_x }}"])
        pie = {
            **METRIC_PAYLOAD,
            "name": "ping_pie",
            "metric_type": "pie",
            "query_config": {
                "host_group_ids": ["4"],
                "host_ids": ["1"],
                "item_ids": ["101"],
            },
        }
        metric = (await api.post("/api/metrics", json=pie)).json()
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {
                        "placeholder": "pie_x",
                        "binding_type": "metric",
                        "binding_ref_id": metric["id"],
                    }
                ]
            },
        )
        assert resp.status_code == 422
        assert "独占一段" in resp.text

    async def test_scalar_metric_can_bind_to_mixed_placeholder(
        self, api: AsyncClient
    ) -> None:
        """混排只限制图表类型；数值指标照常可绑。"""
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["本期 {{ num_x }} Mbps"])
        metric = (
            await api.post("/api/metrics", json={**METRIC_PAYLOAD, "name": "num_x_m"})
        ).json()
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {
                        "placeholder": "num_x",
                        "binding_type": "metric",
                        "binding_ref_id": metric["id"],
                    }
                ]
            },
        )
        assert resp.status_code == 200, resp.text

    async def test_detail_marks_alone_flags(self, api: AsyncClient) -> None:
        """详情返回每个变量的独占标记（供前端警示）。"""
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["{{ alone_x }}", "混排 {{ mixed_x }}"])
        detail = (await api.get(f"/api/templates/{tpl['id']}")).json()
        by_name = {p["name"]: p["is_alone_in_paragraph"] for p in detail["placeholders"]}
        assert by_name == {"alone_x": True, "mixed_x": False}

    async def test_text_metric_can_bind(self, api: AsyncClient) -> None:
        """text 不在护栏内，允许绑定。"""
        await api.post("/api/datasources", json=DATASOURCE_PAYLOAD)
        tpl = await _upload(api, ["{{ note }}"])
        text_metric = {
            **METRIC_PAYLOAD,
            "name": "summary_note",
            "metric_type": "text",
            "query_config": {"host_group_ids": ["4"], "host_ids": ["1"]},
            "type_config": {"text_config": {"mode": "fixed", "fixed_value": "本月无异常"}},
        }
        metric = (await api.post("/api/metrics", json=text_metric)).json()
        resp = await api.put(
            f"/api/templates/{tpl['id']}/bindings",
            json={
                "bindings": [
                    {
                        "placeholder": "note",
                        "binding_type": "metric",
                        "binding_ref_id": metric["id"],
                    }
                ]
            },
        )
        assert resp.status_code == 200
        assert resp.json()["bindings"][0]["binding_ref_id"] == metric["id"]
