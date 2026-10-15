"""AI 摘要自定义提示词 / 参考范文：绑定持久化 + 试跑预览 + 单段重生成。

覆盖用户确认的三项配套能力：
- 提示词 / 范文随绑定落库并回读（一个 AI 占位符 = 一段独立摘要，互不串味）
- 试跑预览：真实 facts + 拟定配置 -> 即时看到 AI 输出（不生成 docx、不落库）
- 单段重生成：用报告 facts 快照重算该段文字并重渲染 docx（**不重算指标**）

`regenerate_insight` 走 `report_service.AsyncSessionLocal` 自建会话，因此夹具里
把它指向测试库；并让数据源适配器在重生成路径上直接抛错，以证明"没重算指标"。
"""

import io
from collections.abc import AsyncGenerator
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
import pytest
import pytest_asyncio
import respx
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import (
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)

import app.api.deps as deps_module
from app.api.deps import get_current_user
from app.config import get_settings
from app.database import get_db
from app.main import app
from app.models import Base
from app.models.report import Report, ReportAudit
from app.models.user import User
from app.schemas.metric import MetricFact
from app.services import report_service
from app.services.adapters.base import DataSourceAdapter
from app.services.preview_service import AI_SUMMARY_PLACEHOLDER

SEP = int(datetime(2026, 9, 1, tzinfo=timezone.utc).timestamp())
PERIOD = "2026-09"
AI_BASE = "http://ai.test/v1"
CHAT_URL = f"{AI_BASE}/chat/completions"

DATASOURCE_PAYLOAD = {
    "name": "生产 Zabbix",
    "type": "zabbix",
    "config": {
        "url": "http://zabbix.example.com",
        "auth_type": "bearer",
        "timeout": 10,
        "verify_ssl": False,
    },
    "secret": "zabbix-token-xyz7890",
}

METRIC_PAYLOAD = {
    "metric_type": "scalar",
    "name": "cpu_usage",
    "data_source_id": 1,
    "query_config": {"host_group_ids": ["4"], "host_ids": ["1"], "item_ids": ["101"]},
    "calc_formula": {"agg": "avg", "compare": "none"},
    "unit": "%",
    "period_type": "monthly",
}


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(path: str | Path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


class FakeAdapter(DataSourceAdapter):
    """返回单条 trend（区间内），聚合后 value_avg=30.0。"""

    async def test_connection(self) -> bool:
        return True

    async def list_host_groups(self) -> list[dict[str, Any]]:
        return [{"group_id": 4, "name": "Linux servers"}]

    async def list_hosts(self, group_ids: list[str]) -> list[dict[str, Any]]:
        return [
            {
                "host_id": 1,
                "name": "web-01",
                "groups": [{"group_id": 4, "name": "Linux servers"}],
            }
        ]

    async def list_items(
        self, host_ids: list[str], value_types: list[int] | None = None
    ) -> list[dict[str, Any]]:
        return [
            {
                "item_id": 101,
                "name": "cpu",
                "key": "system.cpu.util",
                "host_id": 1,
                "host": "web-01",
                "value_type": 0,
                "units": "%",
            }
        ]

    async def fetch_latest_value(self, item_id: int) -> Any:
        return None

    async def get_item_meta(self, item_id: int) -> Any:
        return None

    async def fetch_trend_data(
        self, item_ids: list[int], time_from: int, time_till: int
    ) -> list[dict[str, Any]]:
        return [
            {
                "item_id": 101,
                "clock": SEP + 3600,
                "num": 60,
                "value_min": 30.0,
                "value_avg": 30.0,
                "value_max": 30.0,
            }
        ]


@pytest.fixture
def ai_settings(monkeypatch: pytest.MonkeyPatch) -> Any:
    settings = get_settings()
    monkeypatch.setattr(settings, "ai_base_url", AI_BASE)
    monkeypatch.setattr(settings, "ai_api_key", "test-ai-key")
    monkeypatch.setattr(settings, "ai_model", "test-model")
    return settings


def chat_response(content: str) -> httpx.Response:
    return httpx.Response(
        200,
        json={
            "id": "chatcmpl-1",
            "object": "chat.completion",
            "created": 1,
            "model": "test-model",
            "choices": [
                {
                    "index": 0,
                    "message": {"role": "assistant", "content": content},
                    "finish_reason": "stop",
                }
            ],
        },
    )


@pytest_asyncio.fixture
async def ctx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fake_user: User
) -> AsyncGenerator[dict[str, Any], None]:
    monkeypatch.setattr(get_settings(), "storage_dir", str(tmp_path / "storage"))
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(
        engine, class_=AsyncSession, expire_on_commit=False
    )

    async def _get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = _get_db
    app.dependency_overrides[get_current_user] = lambda: fake_user
    # 报告服务自建会话，指向同一测试库（否则重生成查不到报告）
    monkeypatch.setattr(report_service, "AsyncSessionLocal", session_factory)
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield {
            "client": client,
            "factory": session_factory,
            "tmp": tmp_path,
            "user": fake_user,
        }
    app.dependency_overrides.clear()
    await engine.dispose()


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------
async def _setup_metric(client: AsyncClient) -> int:
    await client.post("/api/datasources", json=DATASOURCE_PAYLOAD)
    resp = await client.post("/api/metrics", json=METRIC_PAYLOAD)
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


async def _upload(
    client: AsyncClient, paragraphs: list[str], name: str = "月报模板"
) -> dict:
    resp = await client.post(
        "/api/templates",
        data={"name": name},
        files={
            "file": ("tpl.docx", make_docx_bytes(paragraphs), "application/octet-stream")
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


async def _put_bindings(
    client: AsyncClient, tpl_id: int, bindings: list[dict]
) -> dict:
    resp = await client.put(
        f"/api/templates/{tpl_id}/bindings", json={"bindings": bindings}
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


async def _current_version_id(client: AsyncClient, tpl_id: int) -> int:
    detail = (await client.get(f"/api/templates/{tpl_id}")).json()
    assert detail["current_version"] is not None
    return detail["current_version"]["id"]


async def _bindings_by_ph(client: AsyncClient, tpl_id: int) -> dict[str, dict]:
    resp = await client.get(f"/api/templates/{tpl_id}")
    assert resp.status_code == 200, resp.text
    return {b["placeholder"]: b for b in resp.json()["bindings"]}


def _fact(metric_id: int, value: float = 30.0) -> MetricFact:
    return MetricFact(
        metric_id=metric_id,
        name="cpu_usage",
        value=value,
        unit="%",
        source="生产 Zabbix",
        object="Linux servers",
        formula="avg(cpu_usage)",
        period=PERIOD,
    )


async def _seed_report(
    factory: async_sessionmaker[AsyncSession],
    *,
    tmp_path: Path,
    template_id: int,
    version_id: int,
    facts: list[MetricFact],
    insights: dict[str, Any] | None = None,
    old_text: str = "旧摘要文本",
) -> int:
    """直接落一条"已生成"的报告（含 facts 快照与 docx 文件）。"""
    old_file = tmp_path / "old_report.docx"
    old_file.write_bytes(make_docx_bytes([old_text]))
    async with factory() as session:
        report = Report(
            template_id=template_id,
            version_id=version_id,
            source="manual",
            period=PERIOD,
            status="pending_confirmation",
            facts_json=[f.model_dump() for f in facts],
            insights_json=insights
            or {
                "summary": {
                    "text": old_text,
                    "unverified_numbers": [],
                    "degraded": False,
                }
            },
            recipient_emails=[],
            file_path=str(old_file),
        )
        session.add(report)
        await session.commit()
        await session.refresh(report)
        return report.id


async def _reload_report(
    factory: async_sessionmaker[AsyncSession], report_id: int
) -> Report:
    async with factory() as session:
        report = await session.get(Report, report_id)
        assert report is not None
        session.expunge(report)
        return report


# --------------------------------------------------------------------------
# 1. 绑定持久化（提示词 / 范文随绑定落库并回读）
# --------------------------------------------------------------------------
class TestBindingAIConfigPersisted:
    async def test_insight_config_persisted_and_read_back(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["CPU: {{ cpu }}", "{{ summary }}"])

        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {
                    "placeholder": "summary",
                    "binding_type": "insight",
                    "ai_config": {
                        "prompt": "只总结 CPU 相关指标，语气简洁",
                        "sample_text": "上期 CPU 为 88.8%，整体平稳。",
                        "max_chars": 300,
                        "temperature": 0.4,
                    },
                },
            ],
        )

        saved = await _bindings_by_ph(client, tpl["id"])
        insight = saved["summary"]
        # 槽位 key 缺省取 placeholder 名 -> 一个 AI 占位符 = 一段独立摘要
        assert insight["binding_key"] == "summary"
        assert insight["ai_config"] == {
            "prompt": "只总结 CPU 相关指标，语气简洁",
            "sample_text": "上期 CPU 为 88.8%，整体平稳。",
            "max_chars": 300,
            "temperature": 0.4,
        }
        # metric 绑定不带 AI 配置，且 binding_key 被清空
        assert saved["cpu"]["ai_config"] is None
        assert saved["cpu"]["binding_key"] is None

    async def test_blank_text_saved_as_no_config(self, ctx: dict) -> None:
        """留空 = 不配置（不是清空为 "" ），下次生成走默认行为。"""
        client: AsyncClient = ctx["client"]
        await _setup_metric(client)
        tpl = await _upload(client, ["{{ summary }}"])

        await _put_bindings(
            client,
            tpl["id"],
            [
                {
                    "placeholder": "summary",
                    "binding_type": "insight",
                    "ai_config": {"prompt": "   ", "sample_text": "\n"},
                }
            ],
        )
        saved = await _bindings_by_ph(client, tpl["id"])
        assert saved["summary"]["ai_config"] is None

    async def test_resave_without_ai_config_clears_it(self, ctx: dict) -> None:
        """绑定是全量替换语义：再次保存时未带 ai_config = 清掉旧配置。"""
        client: AsyncClient = ctx["client"]
        await _setup_metric(client)
        tpl = await _upload(client, ["{{ summary }}"])

        await _put_bindings(
            client,
            tpl["id"],
            [
                {
                    "placeholder": "summary",
                    "binding_type": "insight",
                    "ai_config": {"prompt": "只写一句话"},
                }
            ],
        )
        assert (await _bindings_by_ph(client, tpl["id"]))["summary"]["ai_config"] == {
            "prompt": "只写一句话"
        }

        await _put_bindings(
            client, tpl["id"], [{"placeholder": "summary", "binding_type": "insight"}]
        )
        assert (await _bindings_by_ph(client, tpl["id"]))["summary"]["ai_config"] is None

    async def test_multiple_insight_slots_keep_independent_config(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        await _setup_metric(client)
        tpl = await _upload(
            client, ["{{ overview }}", "{{ risk_detail }}"]
        )

        await _put_bindings(
            client,
            tpl["id"],
            [
                {
                    "placeholder": "overview",
                    "binding_type": "insight",
                    "ai_config": {"prompt": "全局概览", "max_chars": 200},
                },
                {
                    "placeholder": "risk_detail",
                    "binding_type": "insight",
                    # 显式指定槽位 key（允许被覆盖）
                    "binding_key": "risk",
                    "ai_config": {
                        "prompt": "只写风险项",
                        "sample_text": "上月发生过一次磁盘告警。",
                    },
                },
            ],
        )

        saved = await _bindings_by_ph(client, tpl["id"])
        assert saved["overview"]["binding_key"] == "overview"
        assert saved["overview"]["ai_config"] == {"prompt": "全局概览", "max_chars": 200}
        assert saved["risk_detail"]["binding_key"] == "risk"
        assert saved["risk_detail"]["ai_config"]["sample_text"] == "上月发生过一次磁盘告警。"


# --------------------------------------------------------------------------
# 2. 试跑预览（真实 facts + 拟定配置 -> 立即出稿）
# --------------------------------------------------------------------------
class TestInsightPreviewApi:
    @respx.mock
    async def test_preview_returns_text_with_prompt_and_sample(
        self, ctx: dict, ai_settings: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        client: AsyncClient = ctx["client"]
        monkeypatch.setattr(
            deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None)
        )
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["CPU: {{ cpu }}", "{{ summary }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {"placeholder": "summary", "binding_type": "insight"},
            ],
        )

        route = respx.post(CHAT_URL).mock(
            return_value=chat_response("CPU 使用率 30%，运行平稳。")
        )

        resp = await client.post(
            f"/api/templates/{tpl['id']}/insight-preview",
            json={
                "period": PERIOD,
                "placeholder": "summary",
                "ai_config": {
                    "prompt": "只写一句话",
                    "sample_text": "上期 CPU 为 88.8%，整体平稳。",
                },
            },
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["text"] == "CPU 使用率 30%，运行平稳。"
        assert body["degraded"] is False
        assert body["unverified_numbers"] == []
        assert body["sample_leak"] == []
        # 让用户知道 AI 看到了多少数据
        assert body["metric_count"] == 1

        sent = route.calls.last.request.content.decode()
        assert "只写一句话" in sent
        assert "上期 CPU 为 88.8%" in sent

    async def test_preview_degrades_without_api_key(
        self, ctx: dict, ai_settings: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        client: AsyncClient = ctx["client"]
        monkeypatch.setattr(
            deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None)
        )
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["{{ cpu }}", "{{ summary }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {"placeholder": "summary", "binding_type": "insight"},
            ],
        )
        monkeypatch.setattr(get_settings(), "ai_api_key", "")

        resp = await client.post(
            f"/api/templates/{tpl['id']}/insight-preview",
            json={"period": PERIOD, "ai_config": {"prompt": "x"}},
        )
        # AI 不可用不算错误：降级为占位文本，由 degraded 标识
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["degraded"] is True
        assert body["text"] == AI_SUMMARY_PLACEHOLDER

    async def test_preview_rejects_unbound_template(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        tpl = await _upload(client, ["{{ nobody }}"])
        resp = await client.post(
            f"/api/templates/{tpl['id']}/insight-preview",
            json={"period": PERIOD, "ai_config": {}},
        )
        assert resp.status_code == 422
        assert "nobody" in resp.text

    async def test_preview_unknown_template_404(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        resp = await client.post(
            "/api/templates/9999/insight-preview",
            json={"period": PERIOD, "ai_config": {}},
        )
        assert resp.status_code == 404

    async def test_preview_requires_auth(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        app.dependency_overrides.pop(get_current_user, None)
        try:
            resp = await client.post(
                "/api/templates/1/insight-preview",
                json={"period": PERIOD, "ai_config": {}},
            )
            assert resp.status_code == 401
        finally:
            app.dependency_overrides[get_current_user] = lambda: ctx["user"]


# --------------------------------------------------------------------------
# 3. 单段摘要重生成（用快照重算该段 + 重渲染 docx）
# --------------------------------------------------------------------------
class TestRegenerateInsightApi:
    @respx.mock
    async def test_regenerate_uses_snapshot_and_rerenders(
        self, ctx: dict, ai_settings: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        client: AsyncClient = ctx["client"]
        factory: async_sessionmaker[AsyncSession] = ctx["factory"]
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["CPU: {{ cpu }}", "{{ summary }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {
                    "placeholder": "summary",
                    "binding_type": "insight",
                    "ai_config": {"prompt": "只写一句话"},
                },
            ],
        )
        version_id = await _current_version_id(client, tpl["id"])
        report_id = await _seed_report(
            factory,
            tmp_path=ctx["tmp"],
            template_id=tpl["id"],
            version_id=version_id,
            facts=[_fact(metric_id)],
        )

        # 重生成不得重算指标：适配器一旦被调用即失败
        def _boom(*args: Any, **kwargs: Any) -> Any:
            raise AssertionError("重生成不应重算指标（不得加载数据源适配器）")

        monkeypatch.setattr(deps_module, "build_adapter", _boom)
        route = respx.post(CHAT_URL).mock(
            return_value=chat_response("CPU 使用率 30%，运行平稳。")
        )

        resp = await client.post(
            f"/api/reports/{report_id}/insights/summary/regenerate", json={}
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["slot"] == "summary"
        assert body["text"] == "CPU 使用率 30%，运行平稳。"
        assert body["degraded"] is False
        assert route.call_count == 1

        # 模板里配置的提示词被实际使用
        assert "只写一句话" in route.calls.last.request.content.decode()

        report = await _reload_report(factory, report_id)
        assert report.insights_json["summary"]["text"] == "CPU 使用率 30%，运行平稳。"
        # facts 快照原样保留（没有被重算覆盖）
        assert len(report.facts_json) == 1
        assert report.facts_json[0]["value"] == 30.0
        # docx 被就地重渲染，且用的是新文字
        assert report.file_path is not None
        assert report.file_path.endswith(f"report_{report_id}_{PERIOD}.docx")
        rendered = docx_text(report.file_path)
        assert "CPU 使用率 30%，运行平稳。" in rendered
        assert "旧摘要文本" not in rendered

        async with factory() as session:
            audits = (
                await session.scalars(
                    select(ReportAudit).where(ReportAudit.report_id == report_id)
                )
            ).all()
        assert [a.action for a in audits] == ["regenerate_insight"]

    @respx.mock
    async def test_regenerate_accepts_temporary_config_override(
        self, ctx: dict, ai_settings: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """页面改完提示词可先试后存：覆盖配置不落库。"""
        client: AsyncClient = ctx["client"]
        factory: async_sessionmaker[AsyncSession] = ctx["factory"]
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["{{ cpu }}", "{{ summary }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {"placeholder": "summary", "binding_type": "insight"},
            ],
        )
        version_id = await _current_version_id(client, tpl["id"])
        report_id = await _seed_report(
            factory,
            tmp_path=ctx["tmp"],
            template_id=tpl["id"],
            version_id=version_id,
            facts=[_fact(metric_id)],
        )

        monkeypatch.setattr(
            deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None)
        )
        route = respx.post(CHAT_URL).mock(
            return_value=chat_response("CPU 使用率 30%。")
        )

        resp = await client.post(
            f"/api/reports/{report_id}/insights/summary/regenerate",
            json={"ai_config": {"prompt": "临时要求：只写一句"}},
        )
        assert resp.status_code == 200, resp.text
        # 覆盖的提示词进了 user message
        assert "临时要求：只写一句" in route.calls.last.request.content.decode()
        # 但模板绑定没有被改写
        saved = await _bindings_by_ph(client, tpl["id"])
        assert saved["summary"]["ai_config"] is None

    @respx.mock
    async def test_regenerate_migrates_legacy_slot_key(
        self, ctx: dict, ai_settings: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """兼容升级前固定槽位 key：旧报告的 "summary" 会被迁移到当前绑定名。

        否则渲染时按新槽位名取不到文本，这一段会被渲染成占位符。
        """
        client: AsyncClient = ctx["client"]
        factory: async_sessionmaker[AsyncSession] = ctx["factory"]
        metric_id = await _setup_metric(client)
        tpl = await _upload(client, ["{{ cpu }}", "{{ ai_note }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "cpu", "binding_type": "metric", "binding_ref_id": metric_id},
                {"placeholder": "ai_note", "binding_type": "insight"},
            ],
        )
        version_id = await _current_version_id(client, tpl["id"])
        # 历史报告用的是旧 key "summary"
        report_id = await _seed_report(
            factory,
            tmp_path=ctx["tmp"],
            template_id=tpl["id"],
            version_id=version_id,
            facts=[_fact(metric_id)],
            insights={
                "summary": {
                    "text": "旧摘要文本",
                    "unverified_numbers": [],
                    "degraded": False,
                }
            },
        )

        monkeypatch.setattr(
            deps_module, "build_adapter", lambda *a, **kw: FakeAdapter({}, None)
        )
        respx.post(CHAT_URL).mock(return_value=chat_response("CPU 使用率 30%。"))

        resp = await client.post(
            f"/api/reports/{report_id}/insights/summary/regenerate", json={}
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["slot"] == "ai_note"

        report = await _reload_report(factory, report_id)
        assert set(report.insights_json.keys()) == {"ai_note"}
        rendered = docx_text(report.file_path or "")
        assert "CPU 使用率 30%。" in rendered
        assert AI_SUMMARY_PLACEHOLDER not in rendered

    async def test_regenerate_without_snapshot_422(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        factory: async_sessionmaker[AsyncSession] = ctx["factory"]
        tpl = await _upload(client, ["{{ summary }}"])
        await _put_bindings(
            client, tpl["id"], [{"placeholder": "summary", "binding_type": "insight"}]
        )
        version_id = await _current_version_id(client, tpl["id"])
        report_id = await _seed_report(
            factory,
            tmp_path=ctx["tmp"],
            template_id=tpl["id"],
            version_id=version_id,
            facts=[],
        )
        resp = await client.post(
            f"/api/reports/{report_id}/insights/summary/regenerate", json={}
        )
        assert resp.status_code == 422
        assert "事实数据快照" in resp.text

    async def test_regenerate_unknown_slot_422(self, ctx: dict) -> None:
        """多槽位模板必须精确匹配槽位名，不允许乱改一段。"""
        client: AsyncClient = ctx["client"]
        factory: async_sessionmaker[AsyncSession] = ctx["factory"]
        tpl = await _upload(client, ["{{ s1 }}", "{{ s2 }}"])
        await _put_bindings(
            client,
            tpl["id"],
            [
                {"placeholder": "s1", "binding_type": "insight"},
                {"placeholder": "s2", "binding_type": "insight"},
            ],
        )
        version_id = await _current_version_id(client, tpl["id"])
        report_id = await _seed_report(
            factory,
            tmp_path=ctx["tmp"],
            template_id=tpl["id"],
            version_id=version_id,
            facts=[_fact(1)],
            insights={
                "s1": {"text": "旧文本", "unverified_numbers": [], "degraded": False},
                "s2": {"text": "旧文本", "unverified_numbers": [], "degraded": False},
            },
        )
        resp = await client.post(
            f"/api/reports/{report_id}/insights/nope/regenerate", json={}
        )
        assert resp.status_code == 422
        assert "nope" in resp.text

    async def test_regenerate_unknown_report_404(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        resp = await client.post("/api/reports/9999/insights/summary/regenerate", json={})
        assert resp.status_code == 404

    async def test_regenerate_requires_auth(self, ctx: dict) -> None:
        client: AsyncClient = ctx["client"]
        app.dependency_overrides.pop(get_current_user, None)
        try:
            resp = await client.post(
                "/api/reports/1/insights/summary/regenerate", json={}
            )
            assert resp.status_code == 401
        finally:
            app.dependency_overrides[get_current_user] = lambda: ctx["user"]