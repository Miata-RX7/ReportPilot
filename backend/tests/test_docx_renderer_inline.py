"""docx_renderer：context 含 InlineImage 时的渲染与产物校验。

固化 spike 结论：
- InlineImage 与 SandboxedEnvironment 兼容（无需放宽沙箱）
- 独占段落 -> 图片插入，且变量字面量消失
- 混排 -> 图片**也会**插入，但周围文字保留（版式不可控 -> 由绑定校验拦截）
- BytesIO 与文件路径两种图片来源都可用
"""

import io
import zipfile

import pytest
from docx import Document
from docx.shared import Mm
from docxtpl import InlineImage

from app.services.docx_renderer import (
    RenderError,
    check_context,
    open_template,
    render_docx,
    template_variables,
)

def make_png(rgb: tuple[int, int, int]) -> bytes:
    """指定纯色的 1x1 合法 PNG。

    多图测试需要**内容不同**的图片：python-docx 按图片内容哈希去重，
    相同字节的两张图只会产生一个 word/media 条目。
    """
    import struct
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = bytes([0, rgb[0], rgb[1], rgb[2]])
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


PNG = make_png((255, 0, 0))
PNG2 = make_png((0, 0, 255))


def build(paragraphs: list[str], path) -> None:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def media_files(docx_path) -> list[str]:
    with zipfile.ZipFile(docx_path) as zf:
        return [n for n in zf.namelist() if n.startswith("word/media/")]


def paragraph_texts(docx_path) -> list[str]:
    return [p.text for p in Document(str(docx_path)).paragraphs]


class TestInlineImageRender:
    def test_alone_paragraph_embeds_image(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        build(["{{ tag }}"], src)
        tpl = open_template(src)
        context = {"tag": InlineImage(tpl, io.BytesIO(PNG), width=Mm(140))}
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 1
        # 变量字面量不应残留
        assert paragraph_texts(out) == [""]

    def test_file_path_source_also_works(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        png = tmp_path / "c.png"
        png.write_bytes(PNG)
        build(["{{ tag }}"], src)
        tpl = open_template(src)
        context = {"tag": InlineImage(tpl, str(png), width=Mm(140))}
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 1

    def test_mixed_paragraph_also_embeds_but_keeps_text(self, tmp_path) -> None:
        """固化 spike：混排也会插图（因此需要绑定校验拦截，而不是依赖 docxtpl 报错）。"""
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        build(["流量：{{ tag }}"], src)
        tpl = open_template(src)
        context = {"tag": InlineImage(tpl, io.BytesIO(PNG), width=Mm(140))}
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 1
        assert paragraph_texts(out) == ["流量："]

    def test_table_cell_image(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{ tag }}"
        doc.save(str(src))
        tpl = open_template(src)
        context = {"tag": InlineImage(tpl, io.BytesIO(PNG), width=Mm(100))}
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 1

    def test_multiple_images(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        build(["{{ a }}", "{{ b }}"], src)
        tpl = open_template(src)
        context = {
            "a": InlineImage(tpl, io.BytesIO(PNG), width=Mm(140)),
            "b": InlineImage(tpl, io.BytesIO(PNG2), width=Mm(100)),
        }
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 2

    def test_identical_images_deduplicated(self, tmp_path) -> None:
        """python-docx 按内容去重：相同图片只产生一个 media 条目（预期行为）。"""
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        build(["{{ a }}", "{{ b }}"], src)
        tpl = open_template(src)
        context = {
            "a": InlineImage(tpl, io.BytesIO(PNG), width=Mm(140)),
            "b": InlineImage(tpl, io.BytesIO(PNG), width=Mm(100)),
        }
        render_docx(src, context, out, tpl=tpl)
        assert len(media_files(out)) == 1

    def test_without_tpl_argument_still_works(self, tmp_path) -> None:
        """不传 tpl 时（纯文本模板的既有调用方式）行为不变。"""
        src = tmp_path / "t.docx"
        out = tmp_path / "out.docx"
        build(["hello {{ name }}"], src)
        render_docx(src, {"name": "world"}, out)
        assert paragraph_texts(out) == ["hello world"]


class TestContextValidation:
    def test_missing_variable_raises(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        build(["{{ a }} {{ b }}"], src)
        with pytest.raises(RenderError, match="缺失"):
            render_docx(src, {"a": "1"}, tmp_path / "o.docx")

    def test_extra_variable_raises(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        build(["{{ a }}"], src)
        with pytest.raises(RenderError, match="多余"):
            render_docx(src, {"a": "1", "b": "2"}, tmp_path / "o.docx")

    def test_check_context_direct(self) -> None:
        with pytest.raises(RenderError):
            check_context({"a", "b"}, {"a": 1})

    def test_template_variables(self, tmp_path) -> None:
        src = tmp_path / "t.docx"
        build(["{{ a }}{{ b }}"], src)
        assert template_variables(open_template(src)) == {"a", "b"}
