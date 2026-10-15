"""段落独占判定（is_alone_in_paragraph）单测。

独占 = 去掉首尾空白后整段恰好是一个 `{{ var }}`。
这是 docxtpl InlineImage 的前提：混排时图片会被塞进文字中间，版式不可控。
"""

import io as _io

from docx import Document

from app.services.docx_parser import parse_template


def build_docx(paragraphs: list[str], path) -> None:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def alone_map(paragraphs: list[str], tmp_path) -> dict[str, bool]:
    p = tmp_path / "t.docx"
    build_docx(paragraphs, p)
    return parse_template(p).alone


class TestAloneDetection:
    def test_alone_paragraph(self, tmp_path) -> None:
        assert alone_map(["{{ v }}"], tmp_path) == {"v": True}

    def test_mixed_with_prefix(self, tmp_path) -> None:
        assert alone_map(["流量：{{ v }}"], tmp_path) == {"v": False}

    def test_mixed_with_suffix(self, tmp_path) -> None:
        assert alone_map(["{{ v }} bps"], tmp_path) == {"v": False}

    def test_surrounding_whitespace_still_alone(self, tmp_path) -> None:
        """实测 docxtpl 对首尾空白容忍，故判定为独占。"""
        assert alone_map(["  {{ v }}  "], tmp_path) == {"v": True}

    def test_whitespace_inside_braces_alone(self, tmp_path) -> None:
        assert alone_map(["{{v}}"], tmp_path) == {"v": True}
        assert alone_map(["{{ v }}"], tmp_path) == {"v": True}

    def test_two_vars_in_one_paragraph_both_not_alone(self, tmp_path) -> None:
        assert alone_map(["{{ a }}{{ b }}"], tmp_path) == {"a": False, "b": False}

    def test_repeated_occurrence_uses_and(self, tmp_path) -> None:
        """多次出现取 AND：一次独占 + 一次混排 -> 不独占。"""
        assert alone_map(["{{ v }}", "说明 {{ v }}"], tmp_path) == {"v": False}

    def test_repeated_all_alone_stays_alone(self, tmp_path) -> None:
        assert alone_map(["{{ v }}", "{{ v }}"], tmp_path) == {"v": True}

    def test_mixed_order_preserved(self, tmp_path) -> None:
        p = tmp_path / "t.docx"
        build_docx(["{{ z }}", "x {{ a }}", "{{ m }} tail"], p)
        parsed = parse_template(p)
        assert parsed.names == ["z", "a", "m"]
        assert parsed.alone == {"z": True, "a": False, "m": False}

    def test_table_cell_alone(self, tmp_path) -> None:
        doc = Document()
        doc.add_paragraph("head")
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "{{ tbl }}"
        p = tmp_path / "t.docx"
        doc.save(str(p))
        assert parse_template(p).alone == {"tbl": True}

    def test_table_cell_mixed(self, tmp_path) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "值 {{ tbl }}"
        p = tmp_path / "t.docx"
        doc.save(str(p))
        assert parse_template(p).alone == {"tbl": False}

    def test_no_variables(self, tmp_path) -> None:
        p = tmp_path / "t.docx"
        build_docx(["纯文本", "没有变量"], p)
        parsed = parse_template(p)
        assert parsed.names == []
        assert parsed.alone == {}

    def test_split_run_variable_alone(self, tmp_path) -> None:
        """Word 会把 {{ v }} 拆进多个 run；拼接后仍应判定为独占。"""
        doc = Document()
        para = doc.add_paragraph()
        for piece in ("{{", " v", " }}"):
            para.add_run(piece)
        p = tmp_path / "t.docx"
        doc.save(str(p))
        assert parse_template(p).alone == {"v": True}

    def test_as_infos_shape(self, tmp_path) -> None:
        p = tmp_path / "t.docx"
        build_docx(["{{ a }}", "b {{ b }}"], p)
        infos = parse_template(p).as_infos()
        assert [(i.name, i.is_alone_in_paragraph) for i in infos] == [
            ("a", True),
            ("b", False),
        ]


class TestSyntaxStillRejected:
    def test_illegal_fragment_raises(self, tmp_path) -> None:
        from app.services.docx_parser import TemplateSyntaxError

        p = tmp_path / "t.docx"
        build_docx(["{{ a }}", "{% if x %}y{% endif %}"], p)
        try:
            parse_template(p)
        except TemplateSyntaxError:
            return
        raise AssertionError("非法 Jinja 片段应被拒绝")
