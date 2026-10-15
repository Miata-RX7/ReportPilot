"""时间类型端到端：建指标 -> 试算 -> 独占模板绑定 -> 生成报告 -> 校验 docx。"""

import io
import json
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

from docx import Document

from app.services.security import create_token

BASE = "http://localhost:8000"


def api(method, path, token="", body=None, files=None):
    url = BASE + path
    data, headers = None, {}
    if files:
        b = "----rp-t"
        buf = io.BytesIO()
        for name, (fname, content) in files.items():
            buf.write(f"--{b}\r\n".encode())
            buf.write(
                f'Content-Disposition: form-data; name="{name}"; filename="{fname}"\r\n'.encode()
            )
            buf.write(b"Content-Type: application/octet-stream\r\n\r\n")
            buf.write(content)
            buf.write(b"\r\n")
        for k, v in (body or {}).items():
            buf.write(f"--{b}\r\n".encode())
            buf.write(f'Content-Disposition: form-data; name="{k}"\r\n\r\n'.encode())
            buf.write(str(v).encode())
            buf.write(b"\r\n")
        buf.write(f"--{b}--\r\n".encode())
        data = buf.getvalue()
        headers["Content-Type"] = f"multipart/form-data; boundary={b}"
    elif body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=900) as r:
            raw = r.read()
            try:
                return r.status, json.loads(raw.decode())
            except Exception:
                return r.status, raw
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode())
        except Exception:
            return e.code, raw


def token_for(username):
    import asyncio

    from sqlalchemy import select

    from app.database import AsyncSessionLocal
    from app.models.user import User

    async def go():
        async with AsyncSessionLocal() as s:
            u = await s.scalar(select(User).where(User.username == username))
            return create_token(u.id, u.username, u.role, expire_days=1)

    return asyncio.get_event_loop().run_until_complete(go())


def build_template():
    doc = Document()
    doc.add_paragraph("运维月报")
    doc.add_paragraph("报告生成时间：{{ gen_at }}")
    doc.add_paragraph("统计月份：{{ stat_month }}")
    doc.add_paragraph("二、AI 洞察摘要")
    doc.add_paragraph("{{ summary }}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main():
    tok = token_for("ops")

    # 1) 建两个时间指标：上月 / 昨天
    st, last_month = api(
        "POST",
        "/api/metrics",
        tok,
        body={
            "name": "t_last_month",
            "data_source_id": None,
            "unit": "",
            "period_type": "monthly",
            "metric_type": "time",
            "query_config": {},
            "type_config": {
                "time_config": {
                    "format_modules": ["年", "-", "月"],
                    "zero_pad": True,
                    "offset_value": -1,
                    "offset_unit": "month",
                }
            },
        },
    )
    print(f"1) 建上月指标 -> {st} {json.dumps(last_month, ensure_ascii=False)[:120]}")
    assert st == 201, last_month
    st, yesterday = api(
        "POST",
        "/api/metrics",
        tok,
        body={
            "name": "t_yesterday",
            "data_source_id": None,
            "unit": "",
            "period_type": "daily",
            "metric_type": "time",
            "query_config": {},
            "type_config": {
                "time_config": {
                    "format_modules": ["年", "年", "年", "年", "年", "月", "月", "月", "月", "月", "月", "日", "日"],
                    "zero_pad": True,
                    "offset_value": -1,
                    "offset_unit": "day",
                }
            },
        },
    )
    print(f"   建昨日指标 -> {st} (id={yesterday.get('id')})")
    assert st == 201, yesterday

    # 2) 试算
    for label, mid in (("上月", last_month["id"]), ("昨日", yesterday["id"])):
        st, fact = api("POST", f"/api/metrics/{mid}/compute", tok, body={"period": "2026-09"})
        print(
            f"2) 试算{label} -> {st} shape={fact.get('value_shape')} "
            f"value={json.dumps(fact.get('value'), ensure_ascii=False)}"
        )
        print(f"   object={fact.get('object')}")
        assert st == 200 and fact["value_shape"] == "text", fact

    # 3) 列表里 type=data_source_id 应为 null 且不报错
    st, metrics = api("GET", "/api/metrics?limit=50", tok)
    times = [m for m in metrics if m["metric_type"] == "time"]
    print(f"3) 列表里 time 指标 {len(times)} 条，data_source_id={[m['data_source_id'] for m in times]}")
    assert times and all(m["data_source_id"] is None for m in times)

    # 4) 上传独占模板并绑定
    st, tpl = api(
        "POST",
        "/api/templates",
        tok,
        body={"name": f"T时间验证_{int(__import__('time').time())}"},
        files={"file": ("t.docx", build_template())},
    )
    ph = {p["name"]: p["is_alone_in_paragraph"] for p in tpl.get("placeholders", [])}
    print(f"4) 上传模板 -> {st} 独占标记={ph}")
    assert st == 201, tpl
    st, bind = api(
        "PUT",
        f"/api/templates/{tpl['id']}/bindings",
        tok,
        body={
            "bindings": [
                {"placeholder": "gen_at", "binding_type": "metric", "binding_ref_id": yesterday["id"]},
                {"placeholder": "stat_month", "binding_type": "metric", "binding_ref_id": last_month["id"]},
                {"placeholder": "summary", "binding_type": "insight", "binding_key": "summary"},
            ]
        },
    )
    print(f"   绑定 -> {st}")
    assert st == 200, bind

    # 5) 生成报告（时间指标不依赖数据源，即使数据源不可达也应成功）
    st, rep = api(
        "POST",
        "/api/reports",
        tok,
        body={"template_id": tpl["id"], "period": "2026-09", "recipient_emails": ["aiops@livecom.hk"]},
    )
    print(f"5) 生成报告 -> {st} status={rep.get('status')} id={rep.get('id')}")
    if rep.get("status") != "generated":
        print("   last_error:", str(rep.get("last_error"))[:300])
        raise SystemExit("报告未生成成功")

    st, raw = api("GET", f"/api/reports/{rep['id']}/download", tok)
    out = Path("/tmp/report_time.docx")
    out.write_bytes(raw)
    texts = [p.text for p in Document(str(out)).paragraphs]
    print(f"6) docx 段落: {json.dumps(texts, ensure_ascii=False)}")
    assert any("2026-09" in t for t in texts), "报告里没有出现上月时间"

    # 7) 校验 facts_json 里的时间 fact
    st, detail = api("GET", f"/api/reports/{rep['id']}", tok)
    time_facts = [f for f in detail["facts"] if f["value_shape"] == "text"]
    print(f"7) facts 里 text 形态 {len(time_facts)} 条: {[f['value']['text'] for f in time_facts]}")
    assert time_facts

    print("\n端到端通过")
    return last_month["id"], yesterday["id"], tpl["id"], rep["id"]


if __name__ == "__main__":
    main()
